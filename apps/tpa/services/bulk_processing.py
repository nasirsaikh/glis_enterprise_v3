"""Transaction-scoped, versioned TPA worksheets and atomic batch updates."""
import csv
import hashlib
import io
import json
import zipfile
from datetime import date, datetime
from pathlib import Path

from django.core import signing
from django.db import transaction
from django.utils import timezone
from django.utils.crypto import constant_time_compare, salted_hmac

from ..forms import TPAProcessingRowForm
from ..models import MemberAction, MemberPolicyEnrollment, MemberTransaction, TransactionEvent
from .access import can_process_tpa_transaction
from .workflow import validated_tpa_values, _event


COLUMNS = ("transaction_reference", "action_id", "employee_id", "member_name",
           "system_amount", "card_number", "effective_date", "amount",
           "override_reason", "comments", "row_version")
EDITABLE = {"card_number", "effective_date", "amount", "override_reason", "comments"}
MAX_ROWS = 10000
MAX_BYTES = 10 * 1024 * 1024
SIGNING_SALT = "tpa.bulk-processing"


def member_identity(action):
    data = {**(action.submitted_data or {}), **(action.extracted_data or {}), **(action.corrected_data or {})}
    name = data.get("full_name") or " ".join(str(data.get(k) or "") for k in ("first_name", "middle_name", "last_name")).strip()
    return str(data.get("employee_id") or (action.member.employee_id if action.member_id else "")), name or (action.member.full_name if action.member_id else "")


def row_version(action):
    snapshot = [action.pk, action.transaction_id, action.updated_at.isoformat(),
                str(action.calculated_premium), action.card_number,
                str(action.tpa_effective_date or ""), str(action.tpa_premium_amount),
                action.tpa_override_reason, action.processing_message]
    return salted_hmac(SIGNING_SALT, json.dumps(snapshot)).hexdigest()


def _scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return "" if value is None else str(value).strip()


def _csv_safe(value):
    value = _scalar(value)
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def _csv_value(value):
    value = _scalar(value)
    return value[1:] if len(value) > 1 and value[0] == "'" and value[1] in "=+-@\t\r" else value


def build_processing_export(tx, file_format="xlsx"):
    rows = []
    for action in tx.member_actions.select_related("member").order_by("row_number", "pk"):
        employee, name = member_identity(action)
        rows.append([tx.reference, str(action.pk), employee, name, str(action.calculated_premium),
                     action.card_number, (action.tpa_effective_date or tx.effective_date).isoformat(),
                     str(action.tpa_premium_amount if action.tpa_premium_amount is not None else action.calculated_premium),
                     action.tpa_override_reason, action.processing_message, row_version(action)])
    if file_format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(COLUMNS)
        writer.writerows([[_csv_safe(value) for value in row] for row in rows])
        return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TPA Members"
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [27, 14, 20, 30, 20, 24, 20, 20, 42, 48, 44]
    for number, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(number)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="137A55")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            # Identifiers stay text and values beginning '=' are never formulas.
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top")
            cell.protection = Protection(locked=COLUMNS[cell.column - 1] not in EDITABLE)
            if COLUMNS[cell.column - 1] in EDITABLE:
                cell.fill = PatternFill("solid", fgColor="E9F7F0")
    instructions = workbook.create_sheet("Instructions")
    for line in ["Fill the green columns in TPA Members; keep transaction_reference, action_id and row_version unchanged.",
                 "Do not match by row position or name: action_id identifies the member even after sorting.",
                 "Card numbers are text. Preserve leading zeros. Effective date format: YYYY-MM-DD.",
                 "Amount changes need an override_reason. Blank notes preserve existing notes.",
                 "For additions, untouched rows without a card are skipped. Upload only a subset if useful.",
                 "Upload to the same endorsement, review the preview, then Apply. Any invalid row blocks the batch.",
                 "If members changed after downloading, download a fresh workbook. Maximum 10,000 rows per batch."]:
        instructions.append([line])
    instructions.column_dimensions["A"].width = 130
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _read_upload(upload):
    if upload.size > MAX_BYTES:
        raise ValueError("The bulk file exceeds 10 MB.")
    content = upload.read()
    extension = Path(upload.name).suffix.lower()
    if extension == ".csv":
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        headers = next(reader, [])
        raw_rows = reader
    elif extension == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("The expanded workbook exceeds the processing limit.")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = workbook["TPA Members"] if "TPA Members" in workbook.sheetnames else workbook.active
            iterator = sheet.iter_rows()
            headers = [cell.value for cell in next(iterator, [])]
            # Reject formulas instead of silently reading cached formula results.
            raw_rows = []
            for index, row in enumerate(iterator):
                if index >= MAX_ROWS:
                    raise ValueError("Maximum 10,000 rows per import. Split this file into batches.")
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError(f"Row {index + 2} contains a formula. Paste values before uploading.")
                raw_rows.append([cell.value for cell in row])
        finally:
            workbook.close()
    else:
        raise ValueError("Upload the exported XLSX workbook or UTF-8 CSV file.")
    headers = ["_".join(_scalar(item).lower().split()) for item in headers]
    if len(headers) != len(set(headers)):
        raise ValueError("The file has duplicate column headings.")
    if not {"transaction_reference", "action_id", "row_version"}.issubset(headers):
        raise ValueError("Use the downloaded template: transaction_reference, action_id and row_version are required.")
    rows = []
    for number, values in enumerate(raw_rows, 2):
        if number > MAX_ROWS + 1:
            raise ValueError("Maximum 10,000 rows per import. Split this file into batches.")
        if any(value not in (None, "") for value in values):
            if len(values) > len(headers) and any(value not in (None, "") for value in values[len(headers):]):
                raise ValueError(f"Row {number} has more values than column headings.")
            rows.append((number, {key: _csv_value(value) for key, value in zip(headers, values)}))
    if not rows:
        raise ValueError("The uploaded file has no member rows.")
    return rows, hashlib.sha256(content).hexdigest()


def _validate_rows(tx, rows, actions):
    entries, errors, seen, skipped = [], [], set(), 0
    by_id = {str(action.pk): action for action in actions}
    for number, raw in rows:
        identity = raw.get("action_id", "")
        try:
            if raw.get("transaction_reference") != tx.reference:
                raise ValueError("This row belongs to a different endorsement.")
            action = by_id.get(identity)
            if not action:
                raise ValueError("Unknown member row in this endorsement.")
            action.transaction = tx
            if identity in seen:
                raise ValueError("Duplicate action_id in the uploaded file.")
            seen.add(identity)
            if not constant_time_compare(raw.get("row_version", ""), row_version(action)):
                raise ValueError("This member changed after download. Download a fresh member list.")
            employee, name = member_identity(action)
            if (raw.get("employee_id") and raw["employee_id"] != employee) or (raw.get("member_name") and raw["member_name"] != name):
                raise ValueError("The member identity columns were changed. Keep them with their original action_id.")
            values = {"card_number": raw.get("card_number") or action.card_number,
                      "effective_date": raw.get("effective_date") or action.tpa_effective_date or tx.effective_date,
                      "amount": raw.get("amount") or (action.tpa_premium_amount if action.tpa_premium_amount is not None else action.calculated_premium),
                      "override_reason": raw.get("override_reason") or action.tpa_override_reason,
                      "comments": raw.get("comments") or action.processing_message}
            if tx.transaction_type in {tx.Type.MEMBER_ADD, tx.Type.NEW_POLICY_ENROLLMENT} and not values["card_number"]:
                skipped += 1
                continue
            form = TPAProcessingRowForm(values)
            if not form.is_valid():
                raise ValueError("; ".join(f"{field}: {', '.join(messages)}" for field, messages in form.errors.items()))
            cleaned = validated_tpa_values(action, **form.cleaned_data)
            if not tx.policy.start_date <= cleaned["tpa_effective_date"] <= tx.policy.expiry_date:
                raise ValueError("The TPA effective date must be within the policy period.")
            if all(getattr(action, key) == value for key, value in cleaned.items()):
                skipped += 1
                continue
            entries.append({"line": number, "action_id": action.pk, "member_name": name,
                            "row_version": raw["row_version"],
                            "values": {key: _scalar(value) for key, value in cleaned.items()}})
        except (ValueError, TypeError) as exc:
            errors.append({"line": number, "action_id": identity, "message": str(exc)})
    planned = {entry["action_id"]: entry for entry in entries}
    cards = {}
    for action in actions:
        entry = planned.get(action.pk)
        card = entry["values"]["card_number"] if entry else action.card_number
        if card:
            cards.setdefault(card.casefold(), []).append(action.pk)
    occupied = {}
    for enrollment in MemberPolicyEnrollment.objects.filter(policy_id=tx.policy_id).exclude(card_number=""):
        occupied.setdefault(enrollment.card_number.casefold(), set()).add(enrollment.member_id)
    for entry in entries:
        action = by_id[str(entry["action_id"])]
        card = entry["values"]["card_number"].casefold()
        if card and (len(cards.get(card, [])) > 1 or any(member_id != action.member_id for member_id in occupied.get(card, set()))):
            errors.append({"line": entry["line"], "action_id": str(action.pk), "message": "Card number is already assigned to another member in this policy or batch."})
    return {"entries": entries, "errors": errors, "skipped": skipped, "total": len(rows)}


def _require_processing(tx, actor):
    if not can_process_tpa_transaction(actor, tx):
        raise PermissionError("You do not have TPA processing authority.")
    if tx.status != tx.Status.TPA_IN_PROGRESS:
        raise ValueError("Bulk updates are available only while TPA processing is in progress.")


def preview_processing_upload(tx, actor, upload):
    _require_processing(tx, actor)
    rows, digest = _read_upload(upload)
    result = _validate_rows(tx, rows, list(tx.member_actions.select_related("member").order_by("pk")))
    payload = {"transaction_id": tx.pk, "actor_id": actor.pk, "filename": Path(upload.name).name,
               "source_hash": digest, **result}
    token = signing.dumps(payload, salt=SIGNING_SALT, compress=True)
    if len(token) > 1024 * 1024:
        raise ValueError("This preview is too large. Split the file into smaller batches.")
    return {**result, "preview_rows": result["entries"][:20], "error_rows": result["errors"][:50],
            "token": token, "can_apply": bool(result["entries"] and not result["errors"])}


def load_processing_preview(tx, actor, token):
    try:
        data = signing.loads(token, salt=SIGNING_SALT, max_age=1800)
    except signing.BadSignature as exc:
        raise ValueError("The preview expired or changed. Upload the file again.") from exc
    if data.get("transaction_id") != tx.pk or data.get("actor_id") != actor.pk:
        raise PermissionError("This preview belongs to a different user or endorsement.")
    return data


@transaction.atomic
def apply_processing_preview(tx, actor, token):
    tx = MemberTransaction.objects.select_for_update().get(pk=tx.pk)
    _require_processing(tx, actor)
    preview = load_processing_preview(tx, actor, token)
    if preview.get("errors") or not preview.get("entries"):
        raise ValueError("Resolve every import error before applying the batch.")
    actions = list(tx.member_actions.select_for_update().select_related("member").order_by("pk"))
    rows = [(entry["line"], {"transaction_reference": tx.reference, "action_id": str(entry["action_id"]),
            "row_version": entry["row_version"], "card_number": entry["values"]["card_number"],
            "effective_date": entry["values"]["tpa_effective_date"], "amount": entry["values"]["tpa_premium_amount"],
            "override_reason": entry["values"]["tpa_override_reason"], "comments": entry["values"]["processing_message"]})
            for entry in preview["entries"]]
    checked = _validate_rows(tx, rows, actions)
    if checked["errors"]:
        raise ValueError("The batch changed since preview; nothing was saved. " + checked["errors"][0]["message"])
    by_id = {action.pk: action for action in actions}
    updated, events = [], []
    now = timezone.now()
    for entry in checked["entries"]:
        action = by_id[entry["action_id"]]
        values = entry["values"]
        action.card_number = values["card_number"]
        action.tpa_effective_date = date.fromisoformat(values["tpa_effective_date"])
        action.tpa_premium_amount = values["tpa_premium_amount"]
        action.tpa_override_reason = values["tpa_override_reason"]
        action.processing_message = values["processing_message"]
        action.updated_at = now
        updated.append(action)
        events.append(TransactionEvent(transaction=tx, actor=actor, event_type="tpa_member_updated",
            summary=f"Bulk TPA data updated for row {action.row_number or action.pk}.",
            details={"action_id": action.pk, "source_hash": preview["source_hash"], **values}))
    MemberAction.objects.bulk_update(updated, ["card_number", "tpa_effective_date", "tpa_premium_amount",
                                               "tpa_override_reason", "processing_message", "updated_at"], batch_size=500)
    TransactionEvent.objects.bulk_create(events, batch_size=500)
    _event(tx, actor, "tpa_bulk_updated", f"{len(updated)} TPA member rows updated from {preview['filename']}."[:255],
           {"rows_updated": len(updated), "source_hash": preview["source_hash"], "filename": preview["filename"]})
    return len(updated)


def processing_error_csv(errors):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["file_row", "action_id", "error"])
    writer.writerows([[_csv_safe(item[key]) for key in ("line", "action_id", "message")] for item in errors])
    return output.getvalue().encode("utf-8-sig")
