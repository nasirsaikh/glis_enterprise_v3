import csv
import hashlib
import io
from datetime import date, datetime
from pathlib import Path

from django.core.exceptions import ValidationError

from ..models import MemberAction, SourceDocument

MAX_UPLOAD_BYTES = 5 * 1024 * 1024

ALIASES = {
    "member_id": ("member id", "member number", "member no"),
    "tpa_member_id": ("tpa member id", "tpa member number"),
    "card_number": ("card number", "card no"),
    "employee_id": ("employee id", "employee no", "employee number", "staff id"),
    "first_name": ("first name", "given name"),
    "middle_name": ("middle name",),
    "last_name": ("last name", "surname", "family name"),
    "full_name": ("full name", "member name", "name"),
    "date_of_birth": ("date of birth", "dob", "birth date"),
    "gender": ("gender", "sex"),
    "relationship": ("relationship", "relation"),
    "plan_code": ("plan code", "plan", "benefit plan"),
    "national_id": ("national id", "civil id", "id number"),
    "passport_number": ("passport number", "passport", "passport no"),
    "principal_employee_id": (
        "principal employee id",
        "principal employee no",
        "parent employee id",
        "parent employee no",
        "principal emp no",
    ),
    "principal_member_id": (
        "principal member id",
        "parent member id",
        "principal tpa member id",
    ),
}


def _key(value):
    return " ".join(str(value or "").strip().lower().replace("-", " ").replace("_", " ").split())


def _value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else str(value).strip()


def normalize_member_row(raw):
    by_key = {_key(k): v for k, v in raw.items()}
    result = {}
    for target, aliases in ALIASES.items():
        for alias in (target, *aliases):
            normalized_alias = _key(alias)
            if normalized_alias in by_key and by_key[normalized_alias] not in (None, ""):
                result[target] = _value(by_key[normalized_alias])
                break

    full_name = result.pop("full_name", "").strip()
    if full_name and not result.get("first_name"):
        parts = full_name.split()
        result["first_name"] = parts[0]
        result["last_name"] = parts[-1] if len(parts) > 1 else ""
        result["middle_name"] = " ".join(parts[1:-1]) if len(parts) > 2 else ""

    relationship = result.get("relationship", "").upper()
    relation_map = {
        "SELF": "PRINCIPAL", "EMPLOYEE": "PRINCIPAL", "MEMBER": "PRINCIPAL",
        "HUSBAND": "SPOUSE", "WIFE": "SPOUSE", "SON": "CHILD", "DAUGHTER": "CHILD",
    }
    if relationship:
        result["relationship"] = relation_map.get(relationship, relationship)

    gender = result.get("gender", "")
    if gender:
        result["gender"] = {"M": "Male", "F": "Female"}.get(gender.upper(), gender.title())
    dob = result.get("date_of_birth")
    if dob:
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                result["date_of_birth"] = datetime.strptime(dob, fmt).date().isoformat()
                break
            except ValueError:
                continue
    return result


def _read_csv(data):
    return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))


def _read_xlsx(data):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError("Excel upload requires openpyxl from requirements.txt.") from exc

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        return []
    headers = [str(v or "").strip() for v in headers]
    return [
        {headers[i]: value for i, value in enumerate(row)}
        for row in rows
        if any(value not in (None, "") for value in row)
    ]


def import_member_spreadsheet(tx, uploaded_file, actor=None):
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        raise ValidationError("Member upload cannot exceed 5 MB.")

    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ValidationError("Upload a .csv or .xlsx member file.")

    data = uploaded_file.read()
    raw_rows = _read_csv(data) if suffix == ".csv" else _read_xlsx(data)
    if not raw_rows:
        raise ValidationError("The uploaded file does not contain member rows.")

    last_row = tx.member_actions.order_by("-row_number").values_list("row_number", flat=True).first() or 0
    created = []
    for offset, raw in enumerate(raw_rows, start=1):
        normalized = normalize_member_row(raw)
        if not any(normalized.values()):
            continue
        created.append(MemberAction.objects.create(
            transaction=tx,
            action=tx.transaction_type,
            row_number=last_row + offset,
            submitted_data={str(k): _value(v) for k, v in raw.items()},
            corrected_data=normalized,
            extraction_confidence=100,
        ))

    if not created:
        raise ValidationError("No usable member rows were found in the uploaded file.")

    SourceDocument.objects.create(
        transaction=tx,
        original_name=uploaded_file.name,
        document_kind="MEMBER_SPREADSHEET",
        extraction_method="STRUCTURED_IMPORT",
        processed=True,
        extracted_payload={"rows_created": len(created), "format": suffix.lstrip(".")},
        extraction_confidence=100,
        source_hash=hashlib.sha256(data).hexdigest(),
        uploaded_by=actor,
    )
    return created
