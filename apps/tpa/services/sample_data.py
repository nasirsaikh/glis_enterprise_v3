import csv
import io


SAMPLE_COLUMNS = [
    "employee_id",
    "first_name",
    "middle_name",
    "last_name",
    "date_of_birth",
    "gender",
    "relationship",
    "plan_code",
    "national_id",
    "passport_number",
    "principal_employee_id",
    "principal_member_id",
]


def sample_member_rows(tx, include_errors=False):
    plan = tx.policy.plans.filter(is_active=True).order_by("code").first()
    plan_code = plan.code if plan else "PLAN_CODE"

    if include_errors:
        return [
            {
                "employee_id": "ERR-1001",
                "first_name": "Missing",
                "middle_name": "",
                "last_name": "Dob",
                "date_of_birth": "",
                "gender": "Male",
                "relationship": "PRINCIPAL",
                "plan_code": plan_code,
                "national_id": "TEST-ERR-1001",
                "passport_number": "",
                "principal_employee_id": "",
                "principal_member_id": "",
            },
            {
                "employee_id": "ERR-1002",
                "first_name": "Invalid",
                "middle_name": "",
                "last_name": "Plan",
                "date_of_birth": "1992-04-17",
                "gender": "Female",
                "relationship": "PRINCIPAL",
                "plan_code": "INVALID-PLAN",
                "national_id": "TEST-ERR-1002",
                "passport_number": "",
                "principal_employee_id": "",
                "principal_member_id": "",
            },
            {
                "employee_id": "ERR-1003",
                "first_name": "No",
                "middle_name": "",
                "last_name": "Principal",
                "date_of_birth": "2016-11-02",
                "gender": "Female",
                "relationship": "CHILD",
                "plan_code": plan_code,
                "national_id": "TEST-ERR-1003",
                "passport_number": "",
                "principal_employee_id": "",
                "principal_member_id": "",
            },
        ]

    return [
        {
            "employee_id": "SAMPLE-1001",
            "first_name": "Ahmed",
            "middle_name": "Ali",
            "last_name": "Al Harthi",
            "date_of_birth": "1988-05-12",
            "gender": "Male",
            "relationship": "PRINCIPAL",
            "plan_code": plan_code,
            "national_id": "TEST-CID-1001",
            "passport_number": "TEST-P-1001",
            "principal_employee_id": "",
            "principal_member_id": "",
        },
        {
            "employee_id": "SAMPLE-1002",
            "first_name": "Aisha",
            "middle_name": "",
            "last_name": "Al Harthi",
            "date_of_birth": "1990-09-22",
            "gender": "Female",
            "relationship": "SPOUSE",
            "plan_code": plan_code,
            "national_id": "TEST-CID-1002",
            "passport_number": "TEST-P-1002",
            "principal_employee_id": "SAMPLE-1001",
            "principal_member_id": "",
        },
        {
            "employee_id": "SAMPLE-1003",
            "first_name": "Mariam",
            "middle_name": "",
            "last_name": "Al Harthi",
            "date_of_birth": "2018-03-15",
            "gender": "Female",
            "relationship": "CHILD",
            "plan_code": plan_code,
            "national_id": "TEST-CID-1003",
            "passport_number": "",
            "principal_employee_id": "SAMPLE-1001",
            "principal_member_id": "",
        },
    ]


def build_sample_csv(rows):
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=SAMPLE_COLUMNS)
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def build_sample_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Members"
    sheet.append(SAMPLE_COLUMNS)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    for row in rows:
        sheet.append([row.get(column, "") for column in SAMPLE_COLUMNS])

    for index, column in enumerate(SAMPLE_COLUMNS, start=1):
        width = max(len(column) + 2, 14)
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 28)

    info = workbook.create_sheet("Instructions")
    info.append(["Column", "Required", "Notes"])
    instructions = [
        ("employee_id", "Recommended", "Employee/member number. Used to link dependents in the same upload."),
        ("first_name", "Yes", "Required for new enrollment/member addition."),
        ("middle_name", "No", "Optional."),
        ("last_name", "Yes", "Required for new enrollment/member addition."),
        ("date_of_birth", "Yes", "Use YYYY-MM-DD."),
        ("gender", "Yes", "Male or Female."),
        ("relationship", "Yes", "PRINCIPAL, SPOUSE, CHILD or OTHER."),
        ("plan_code", "Yes", "Must be an active benefit plan code under the selected policy."),
        ("national_id", "Optional", "Civil/National ID; duplicate active IDs are rejected."),
        ("passport_number", "Optional", "Passport number."),
        ("principal_employee_id", "Dependent only", "For spouse/child/other, use the principal employee_id. It may refer to a principal row in the same file."),
        ("principal_member_id", "Dependent only", "Alternative: existing TPA principal member ID."),
    ]
    for item in instructions:
        info.append(item)
    for cell in info[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 16
    info.column_dimensions["C"].width = 88

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
