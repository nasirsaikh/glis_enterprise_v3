import csv
import io
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from .models import MemberAction, MemberTransaction, Policy, TPAOrganization
from .services.bulk_processing import (
    apply_processing_preview, build_processing_export, preview_processing_upload,
)


class TPABulkProcessingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.actor = get_user_model().objects.create_superuser("bulk-admin", "bulk@example.test", "test")
        cls.other_actor = get_user_model().objects.create_superuser("bulk-other", "other@example.test", "test")
        cls.scoped = get_user_model().objects.create_user("bulk-scoped", is_staff=True)
        cls.sponsor = TPAOrganization.objects.create(code="BK-SP", name_en="Own sponsor", organization_type="CORPORATE")
        cls.other_sponsor = TPAOrganization.objects.create(code="BK-OTHER", name_en="Other sponsor", organization_type="CORPORATE")
        cls.insurer = TPAOrganization.objects.create(code="BK-IN", name_en="Insurer", organization_type="INSURER")
        cls.scoped.profile.organizations.add(cls.sponsor)
        cls.scoped.user_permissions.add(Permission.objects.get(codename="configure_tpa"))
        cls.policy = Policy.objects.create(sponsor=cls.sponsor, insurance_company=cls.insurer,
            policy_number="BULK-1", start_date=date(2026, 1, 1), expiry_date=date(2026, 12, 31))
        cls.other_policy = Policy.objects.create(sponsor=cls.other_sponsor, insurance_company=cls.insurer,
            policy_number="BULK-2", start_date=date(2026, 1, 1), expiry_date=date(2026, 12, 31))

    def setUp(self):
        self.tx = self.make_transaction(self.policy)
        self.rows = [self.make_action(self.tx, number) for number in (1, 2)]
        self.client.force_login(self.actor)

    def make_transaction(self, policy):
        return MemberTransaction.objects.create(policy=policy, sponsor=policy.sponsor,
            insurer=self.insurer, requester=self.actor, requester_organization=policy.sponsor,
            transaction_type="MEMBER_ADD", status="tpa_in_progress", effective_date=date(2026, 7, 1))

    def make_action(self, tx, number):
        return MemberAction.objects.create(transaction=tx, action="MEMBER_ADD", row_number=number,
            corrected_data={"employee_id":f"00{number}", "first_name":"Sam", "last_name":f"Example{number}"},
            validation_status="VALID", calculated_premium=Decimal("100.000"))

    def csv_rows(self, tx=None):
        content, _ = build_processing_export(tx or self.tx, "csv")
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))

    def upload(self, rows):
        output = io.StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=list(rows[0]))
        writer.writeheader();writer.writerows(rows)
        return SimpleUploadedFile("cards.csv", output.getvalue().encode("utf-8"))

    def filled(self):
        rows = self.csv_rows()
        for index, row in enumerate(rows, 1):
            row["card_number"] = f"000{index}"
        return rows

    def preview(self, rows=None):
        return preview_processing_upload(self.tx, self.actor, self.upload(rows or self.filled()))

    def test_export_uses_text_ids_and_no_formulas(self):
        self.rows[0].processing_message = "=HYPERLINK(\"https://example.test\")"
        self.rows[0].save()
        content, _ = build_processing_export(self.tx)
        workbook = load_workbook(io.BytesIO(content))
        self.assertEqual(workbook.active["C2"].value, "001")
        self.assertEqual(workbook.active["J2"].data_type, "s")
        self.assertEqual(workbook.active["J2"].value, self.rows[0].processing_message)
        self.assertIn("Instructions", workbook.sheetnames)
        workbook.close()

    def test_csv_sorted_rows_match_by_id_and_preserve_leading_zero_cards(self):
        rows = list(reversed(self.filled()))
        preview = self.preview(rows)
        self.assertTrue(preview["can_apply"])
        self.assertEqual(apply_processing_preview(self.tx, self.actor, preview["token"]), 2)
        for index, action in enumerate(self.rows, 1):
            action.refresh_from_db()
            self.assertEqual(action.card_number, f"000{index}")
            self.assertEqual(action.tpa_effective_date, self.tx.effective_date)
            self.assertEqual(action.tpa_premium_amount, Decimal("100.000"))
        self.assertEqual(self.tx.events.filter(event_type="tpa_member_updated").count(), 2)
        self.assertTrue(self.tx.events.filter(event_type="tpa_bulk_updated").exists())

    def test_preview_never_saves_and_can_import_subset(self):
        preview = self.preview(self.filled()[:1])
        self.rows[0].refresh_from_db()
        self.assertEqual(self.rows[0].card_number, "")
        self.assertEqual(apply_processing_preview(self.tx, self.actor, preview["token"]), 1)
        self.rows[1].refresh_from_db()
        self.assertEqual(self.rows[1].card_number, "")

    def test_blank_unassigned_members_are_skipped(self):
        preview = self.preview(self.csv_rows())
        self.assertEqual(preview["skipped"], 2)
        self.assertFalse(preview["can_apply"])

    def test_duplicate_cards_block_the_entire_batch(self):
        rows = self.filled();rows[1]["card_number"] = rows[0]["card_number"]
        preview = self.preview(rows)
        self.assertEqual(len(preview["errors"]), 2)
        with self.assertRaises(ValueError):
            apply_processing_preview(self.tx, self.actor, preview["token"])
        self.assertFalse(self.tx.member_actions.exclude(card_number="").exists())

    def test_duplicate_member_and_identity_changes_are_rejected(self):
        rows = self.filled();rows[1]["action_id"] = rows[0]["action_id"]
        self.assertIn("Duplicate action_id", self.preview(rows)["errors"][0]["message"])
        rows = self.filled();rows[0]["employee_id"] = "wrong-person"
        self.assertIn("identity", self.preview(rows)["errors"][0]["message"])

    def test_other_endorsement_member_id_is_rejected(self):
        other = self.make_transaction(self.other_policy)
        foreign = self.make_action(other, 1)
        rows = self.filled();rows[0]["action_id"] = str(foreign.pk)
        self.assertIn("Unknown member", self.preview(rows)["errors"][0]["message"])
        rows = self.filled();rows[0]["transaction_reference"] = other.reference
        self.assertIn("different endorsement", self.preview(rows)["errors"][0]["message"])

    def test_amount_override_needs_reason_and_bad_dates_are_rejected(self):
        rows = self.filled();rows[0]["amount"] = "150.000"
        self.assertIn("override reason", self.preview(rows)["errors"][0]["message"])
        rows[0]["override_reason"] = "Agreed correction"
        self.assertTrue(self.preview(rows)["can_apply"])
        rows[0]["effective_date"] = "not-a-date"
        self.assertIn("effective_date", self.preview(rows)["errors"][0]["message"])

    def test_stale_row_after_preview_rolls_back_every_member(self):
        preview = self.preview()
        self.rows[1].processing_message = "Changed after preview"
        self.rows[1].save()
        with self.assertRaisesRegex(ValueError, "nothing was saved"):
            apply_processing_preview(self.tx, self.actor, preview["token"])
        self.rows[0].refresh_from_db()
        self.assertEqual(self.rows[0].card_number, "")

    def test_preview_is_user_bound_and_cannot_be_replayed(self):
        preview = self.preview()
        with self.assertRaises(PermissionError):
            apply_processing_preview(self.tx, self.other_actor, preview["token"])
        apply_processing_preview(self.tx, self.actor, preview["token"])
        with self.assertRaises(ValueError):
            apply_processing_preview(self.tx, self.actor, preview["token"])

    def test_tampered_token_and_closed_status_are_rejected(self):
        preview = self.preview()
        with self.assertRaises(ValueError):
            apply_processing_preview(self.tx, self.actor, preview["token"] + "changed")
        self.tx.status = "card_dispatch";self.tx.save()
        with self.assertRaises(ValueError):
            apply_processing_preview(self.tx, self.actor, preview["token"])

    def test_existing_member_endorsements_do_not_issue_new_cards(self):
        self.tx.transaction_type = "MEMBER_DELETE";self.tx.save()
        self.rows[0].card_number = "EXISTING";self.rows[0].save()
        rows = self.csv_rows();rows[0]["card_number"] = "NEW-CARD"
        self.assertIn("cannot issue", self.preview(rows)["errors"][0]["message"])

    def test_blank_notes_preserve_existing_notes(self):
        self.rows[0].processing_message = "Keep this note";self.rows[0].save()
        rows = self.filled();rows[0]["comments"] = ""
        apply_processing_preview(self.tx, self.actor, self.preview(rows)["token"])
        self.rows[0].refresh_from_db()
        self.assertEqual(self.rows[0].processing_message, "Keep this note")

    def test_xlsx_upload_and_formula_rejection(self):
        content, _ = build_processing_export(self.tx)
        workbook = load_workbook(io.BytesIO(content));sheet=workbook.active
        sheet["F2"]="001234";sheet["F3"]="001235"
        output=io.BytesIO();workbook.save(output)
        preview=preview_processing_upload(self.tx,self.actor,SimpleUploadedFile("cards.xlsx",output.getvalue()))
        self.assertTrue(preview["can_apply"])
        sheet["H2"]="=1+1";output=io.BytesIO();workbook.save(output);workbook.close()
        with self.assertRaisesRegex(ValueError,"formula"):
            preview_processing_upload(self.tx,self.actor,SimpleUploadedFile("cards.xlsx",output.getvalue()))

    def test_member_list_is_paginated_and_searches_all_rows(self):
        for number in range(3, 27):self.make_action(self.tx,number)
        response=self.client.get(reverse("tpa:transaction_detail",args=[self.tx.reference]))
        self.assertEqual(len(response.context["processing_page_obj"]),20)
        self.assertEqual(response.content.count(b"data-tpa-notes-dialog"),20)
        self.assertNotContains(response,'id_tpa_999')
        response=self.client.get(reverse("tpa:transaction_detail",args=[self.tx.reference]),{"q":"Example26"})
        self.assertEqual(len(response.context["processing_page_obj"]),1)
        self.assertContains(response,"Example26")
        self.assertContains(response,'data-table-pagination="server"')

    def test_export_and_import_endpoints_respect_organization_scope(self):
        other=self.make_transaction(self.other_policy)
        self.client.force_login(self.scoped)
        self.assertEqual(self.client.get(reverse("tpa:transaction_tpa_export",args=[self.tx.reference])).status_code,200)
        self.assertEqual(self.client.get(reverse("tpa:transaction_tpa_export",args=[other.reference])).status_code,404)
        self.assertEqual(self.client.post(reverse("tpa:transaction_tpa_bulk",args=[other.reference]),{"file":self.upload(self.filled())}).status_code,404)

    def test_htmx_preview_reopens_modal_and_apply_refreshes_workspace(self):
        url=reverse("tpa:transaction_tpa_bulk",args=[self.tx.reference])
        response=self.client.post(url,{"mode":"preview","file":self.upload(self.filled())},HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'data-reopen-modal="tpa-bulk-import"')
        token=response.context["bulk_tpa_preview"]["token"]
        response=self.client.post(url,{"mode":"apply","preview_token":token},HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.context["processing_ready_count"],2)

    def test_a_thousand_members_apply_in_batches(self):
        self.tx.member_actions.all().delete()
        MemberAction.objects.bulk_create([MemberAction(transaction=self.tx,action="MEMBER_ADD",row_number=n,
            corrected_data={"employee_id":str(n),"first_name":"Sam","last_name":str(n)},
            validation_status="VALID",calculated_premium=100) for n in range(1000)])
        rows=self.csv_rows()
        for number,row in enumerate(rows):row["card_number"]=f"CARD-{number:06d}"
        with CaptureQueriesContext(connection) as queries:
            preview=self.preview(rows)
            self.assertEqual(apply_processing_preview(self.tx,self.actor,preview["token"]),1000)
        self.assertLess(len(queries),100)
        self.assertEqual(self.tx.member_actions.exclude(card_number="").count(),1000)
